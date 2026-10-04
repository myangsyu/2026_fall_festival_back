"""부스 정리 엑셀(축제부스_전체정리_vN.xlsx)을 seed_booths용 JSON으로 변환한다.

엑셀은 대협·재원이 관리하는 원본이라 레포에 넣지 않고, 이 커맨드로 뽑은
JSON(apps/booths/data/booths.json)만 커밋한다. 서버는 JSON만 읽으므로
openpyxl은 이 커맨드를 돌리는 로컬에만 있으면 된다.

    pip install openpyxl
    python manage.py convert_booth_xlsx "축제부스_전체정리_v7.xlsx"

시트/열 구성이 바뀌면 HEADER_ROW와 각 시트의 열 이름만 맞춰주면 된다.
"""

import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError

DEFAULT_OUTPUT = Path(__file__).resolve().parents[2] / "data" / "booths.json"

# 모든 데이터 시트는 1~3행이 제목/설명/빈 줄이고 4행이 헤더다.
HEADER_ROW = 4

# 엑셀에서는 이름이 같아 목록에서 구분이 안 되는 부스. DB에는 migration 0005로 이미
# 바뀐 이름이 들어가 있어서, 엑셀 이름 그대로 넣으면 (name, zone)이 달라져 중복 부스가 생긴다.
RENAMED_BOOTHS = {
    ("다회용기 부스", "혜화관"): "다회용기 부스 (혜화관)",
    ("다회용기 부스", "팔정도"): "다회용기 부스 (팔정도)",
}

# 엑셀 category 제안과 다르게 운영에서 확정한 협업(COLLAB)/일반(ETC) 분류.
# (name, zone)은 RENAMED_BOOTHS 적용 후 이름 기준이다.
CATEGORY_OVERRIDES = {
    ("의료인공지능학과", "팔정도"): "COLLAB",
    ("축기단", "팔정도"): "COLLAB",
    ("동국 108리더스", "팔정도"): "COLLAB",
    ("다회용기 부스 (혜화관)", "혜화관"): "COLLAB",
    ("다회용기 부스 (팔정도)", "팔정도"): "COLLAB",
    ("애드러쉬", "팔정도"): "ETC",
    ("오뚜기 진라면 서포터즈 진앤지니", "팔정도"): "ETC",
}

# ── 엑셀(v10) 이후 부스에서 직접 받은 수정사항 ──────────────────────────────
# 엑셀을 다시 변환해도 유지되도록 여기서 덮어쓴다. 키는 (name, zone).

# 부스 필드 덮어쓰기. entrance_fee는 원(인당).
BOOTH_OVERRIDES = {
    ("열린전공학부", "혜화관"): {"entrance_fee": 4000},
    ("광고홍보학과", "혜화관"): {"entrance_fee": 5000},
    ("참사람봉사단", "팔정도"): {
        "subtitle": "주식회사 한잔",
        "description": (
            "📢 주식회사 한잔 신입사원 채용공고\n"
            "모집 부문: 오늘 밤 함께 마실 사람 (경력 무관)\n"
            "근무 조건: 칼퇴 보장, 회식 필참, 야근은 안주로만\n"
            "복리후생: 골뱅이소면, 통삼겹 부추무침, 나가사끼짬뽕 무제한 주문 가능\n"
            "지원 방법: 빈 테이블에 착석 시 즉시 입사 처리"
        ),
        "event_description": (
            "💝 자리값 할인 이벤트 : 생명나눔캠페인 참여시 자리값 5000원 할인\n\n"
            "😈 몬스터 증정 이벤트 : 당신의 야근을 위한 에너지 음료 몬스터 1캔 무료\n\n"
            "👩‍❤️‍👨 미팅 이벤트 : 사원들간의 시너지를 위해 미팅을 잡아드립니다\n\n"
            "📸 우수사원 포토 이벤트 : 메인 메뉴 3개 이상 주문 시 폴라로이드로 사진 촬영"
            " (40명 선착순)"
        ),
    },
    ("공과대학", "원흥관"): {
        "subtitle": "🎃🔥 공대 야간부스 「호박나이트」 🔥🎃",
        "location_detail": "1번(원흥관 주차장)",
        "description": (
            "안녕하세요! 동국대학교 공과대학 학생회 [CONNECT] 입니다.\n"
            "축제 둘째날인 9월 30일! 저희가 야간부스를 하게 되었습니다🎉\n\n"
            "호박나이트에서\n"
            "이 날 만큼은 과제 말고 안주에 집중하세요.\n"
            "그리고 잊지 못할 밤을 즐겨보세요! 🍻"
        ),
        "event_description": (
            "🎃 호박나이트 EVENT 🎃\n\n"
            "원흥관과 신공학관 복도에 붙어있는 포스터를 그냥 지나치지 마세요👀\n"
            "포스터 뒤에 숨겨진 메뉴 쿠폰권이 있다는 사실!\n\n"
            "🔎 원흥관 & 신공학관 복도에 붙어있는 이벤트 포스터를 찾아 숨겨진 메뉴 쿠폰권을"
            " 찾아주세요!\n\n"
            "🎟️ 쿠폰을 찾았다면?\n"
            "9월 30일(수) 오후 6시 30분 전까지 쿠폰을 지참하고\n"
            "📍 원흥관 주차장 공과대학 학생회 야간부스로 오면\n"
            "✨ 쿠폰에 적힌 해당 메뉴를 무료로 서비스! ✨"
        ),
    },
}

# 운영일정 덮어쓰기. dates가 있으면 그 날짜만 남기고, open_at/close_at은 남은 운영일 전체에 적용.
OPERATION_OVERRIDES = {
    ("광고홍보학과", "혜화관"): {"dates": ["2026-09-30"], "close_at": "22:00"},
    ("참사람봉사단", "팔정도"): {"open_at": "18:00"},
    ("공과대학", "원흥관"): {"dates": ["2026-09-30"], "open_at": "18:00"},
    # 10/1 운영 취소
    ("정치외교학전공", "혜화관"): {"dates": ["2026-09-29", "2026-09-30"]},
    ("행정학전공", "혜화관"): {"dates": ["2026-09-29", "2026-09-30"]},
    ("사회복지상담학과", "혜화관"): {"dates": []},
}

# 메뉴 통째로 교체. (메뉴명, 가격)
MENU_OVERRIDES = {
    ("참사람봉사단", "팔정도"): [
        ("SET 메뉴1 (나가사끼짬뽕 + 골뱅이무침, 소면)", 30000),
        ("SET 메뉴2 (나가사끼짬뽕 + 통삼겹 부추무침)", 33000),
        ("SET 메뉴3 (나가사끼짬뽕 + 콘치즈 닭발)", 32000),
        ("골뱅이무침 & 소면", 17000),
        ("통삼겹부추무침", 20000),
        ("나가사끼짬뽕", 17000),
        ("콘치즈닭발", 19000),
    ],
    ("공과대학", "원흥관"): [
        ("SET 1 부대주먹다리 (부대찌개 + 불닭닭다리구이 + 주먹밥)", 36900),
        ("SET 2 부대김치전찌개 (부대찌개 + 김치전)", 34900),
        ("부대찌개 (라면사리 포함)", 18900),
        ("불주먹다리 (순살 불닭닭다리구이 + 주먹밥)", 18900),
        ("장인의 김치지짐 (김치전)", 16900),
        ("아이스 불바나나 (바나나브륄레 + 바닐라아이스크림)", 10900),
        ("콘치즈", 9900),
        ("설탕토마토", 9900),
        ("춤초나쵸 (나쵸 + 치즈소스 + 야채토핑)", 9900),
    ],
}

# 엑셀 원문 오타 수정. 메뉴명 원문 → 수정본.
MENU_NAME_FIXES = {
    "스파르타불닭 (치즈불닭볶음변)": "스파르타불닭 (치즈불닭볶음면)",
}

# 구조물 열 예: "MARKET 21×12m" → placements에 structure/width/depth로 나간다.
STRUCTURE_PATTERN = re.compile(r"^(?P<kind>[A-Z]+)\s+(?P<width>[\d.]+)\s*[×x]\s*(?P<depth>[\d.]+)")


def _sheet_rows(workbook, sheet_name):
    """헤더 첫 줄(개행 앞)을 키로 한 dict 목록. 임시키/날짜가 비어 있는 줄은 건너뛴다."""
    rows = workbook[sheet_name].iter_rows(min_row=HEADER_ROW, values_only=True)
    header = [str(cell).split("\n")[0].strip() if cell else None for cell in next(rows)]
    for row in rows:
        if row[0] is None:
            continue
        yield dict(zip(header, row, strict=False))


def _text(value):
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def _number(value):
    if value is None:
        return None
    number = float(value)
    return int(number) if number.is_integer() else number


def _date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    return str(value)


def _bool(value, formula, label):
    """엑셀 TRUE/FALSE 셀을 bool로. 값 계산 없이 저장된 파일은 캐시 값이 비어 있어서
    수식 원문(=TRUE())으로 판단한다. 둘 다 없으면 추측하지 않고 멈춘다."""
    if isinstance(value, bool):
        return value
    text = str(formula or "").strip().upper().lstrip("=").removesuffix("()")
    if text in ("TRUE", "FALSE"):
        return text == "TRUE"
    raise CommandError(f"{label}: TRUE/FALSE를 읽을 수 없습니다 (값 {value!r}, 원문 {formula!r})")


def _structure(value):
    value = _text(value)
    if not value:
        return {}
    match = STRUCTURE_PATTERN.match(value)
    if not match:
        raise CommandError(f"구조물 표기를 해석할 수 없습니다: {value!r}")
    return {
        "structure": match["kind"],
        "width": _number(match["width"]),
        "depth": _number(match["depth"]),
    }


class Command(BaseCommand):
    help = "부스 정리 엑셀을 seed_booths용 JSON으로 변환한다 (로컬 전용, openpyxl 필요)."

    def add_arguments(self, parser):
        parser.add_argument("xlsx_path", help="축제부스_전체정리_vN.xlsx 경로")
        parser.add_argument("--output", default=str(DEFAULT_OUTPUT), help="출력 JSON 경로")

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError as exc:
            raise CommandError("openpyxl이 필요합니다: pip install openpyxl") from exc

        xlsx_path = Path(options["xlsx_path"])
        workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
        # 수식 원문. 계산 값이 저장되지 않은 셀을 읽을 때만 쓴다.
        formula_booths = {
            row["임시키"]: row for row in _sheet_rows(openpyxl.load_workbook(xlsx_path), "부스")
        }

        placements = defaultdict(list)
        for row in _sheet_rows(workbook, "배치좌표"):
            key = (row["임시키"], _date(row["festival_date"]), row["time_slot"])
            placements[key].append(
                {
                    "unit_no": int(row["천막 번호"]),
                    "map_x": _number(row["map_x"]),
                    "map_y": _number(row["map_y"]),
                    "map_elevation": _number(row["map_elevation"]),
                    "rotation": _number(row["rotation"]),
                    "booth_size": row["booth_size"],
                    **_structure(row["구조물"]),
                }
            )

        operations = defaultdict(list)
        for row in _sheet_rows(workbook, "운영일정"):
            festival_date = _date(row["festival_date"])
            key = (row["임시키"], festival_date, row["time_slot"])
            operations[row["임시키"]].append(
                {
                    "festival_date": festival_date,
                    "time_slot": row["time_slot"],
                    "open_at": row["open_at"],
                    "close_at": row["close_at"],
                    "placements": sorted(placements.pop(key, []), key=lambda p: p["unit_no"]),
                }
            )
        if placements:
            orphans = ", ".join("/".join(key) for key in placements)
            raise CommandError(f"운영일정에 없는 배치좌표가 있습니다: {orphans}")

        # 'DB 반영'이 '포함'인 메뉴만 넣는다. 가격 미수령 행은 price NOT NULL이라 제외.
        menus = defaultdict(list)
        skipped_menus = 0
        for row in _sheet_rows(workbook, "메뉴"):
            if row["DB 반영"] != "포함":
                skipped_menus += 1
                continue
            menu_name = _text(row["메뉴명(name)"])
            menus[row["임시키"]].append(
                {
                    "name": MENU_NAME_FIXES.get(menu_name, menu_name),
                    "price": int(row["price(원)"]),
                }
            )

        booths = []
        skipped_booths = []
        for row in _sheet_rows(workbook, "부스"):
            key = row["임시키"]
            # 좌표가 없으면 지도에 그릴 수 없으므로 좌표를 받을 때까지 넣지 않는다.
            # 화장실은 건물 안이라 좌표 없이 바텀시트 목록에만 보여주므로 예외로 넣는다.
            has_position = row["map_x"] is not None and row["map_y"] is not None
            if not has_position and row["category"] != "TOILET":
                skipped_booths.append(f"{key} {_text(row['name'])}")
                operations.pop(key, None)
                menus.pop(key, None)
                continue

            category = row["category"]
            # 지도 '주류'(ALCOHOL) 칩은 주류 판매 시설만 보여준다. 주·야간 부스(주점 포함)는
            # 전부 '부스' 칩(COLLAB·ETC)으로 모은다.
            if row["place_type"] == "BOOTH" and category == "ALCOHOL":
                category = "ETC"

            name = _text(row["name"])
            zone = _text(row["zone"])
            name = RENAMED_BOOTHS.get((name, zone), name)
            category = CATEGORY_OVERRIDES.get((name, zone), category)

            # v10부터 있는 열. 이전 버전 엑셀에는 없어서 비어 있는 것으로 본다.
            restroom_type = _text(row.get("restroom_type"))
            if restroom_type not in (None, "MALE", "FEMALE", "BOTH"):
                raise CommandError(f"{key} restroom_type 값이 올바르지 않습니다: {restroom_type!r}")

            booth_operations = operations.pop(key, [])
            operation_override = OPERATION_OVERRIDES.get((name, zone), {})
            if "dates" in operation_override:
                booth_operations = [
                    operation
                    for operation in booth_operations
                    if operation["festival_date"] in operation_override["dates"]
                ]
            for operation in booth_operations:
                for field in ("open_at", "close_at"):
                    if field in operation_override:
                        operation[field] = operation_override[field]

            booth_menus = menus.pop(key, [])
            if (name, zone) in MENU_OVERRIDES:
                booth_menus = [
                    {"name": menu_name, "price": price}
                    for menu_name, price in MENU_OVERRIDES[(name, zone)]
                ]

            booth = {
                "key": key,
                "name": name,
                "subtitle": _text(row["subtitle"]),
                "place_type": row["place_type"],
                "category": category,
                "restroom_type": restroom_type,
                "booth_size": row["booth_size"],
                "zone": zone,
                "location_detail": _text(row["location_detail"]),
                "map_x": _number(row["map_x"]),
                "map_y": _number(row["map_y"]),
                "map_elevation": _number(row["map_elevation"]),
                "rotation": _number(row["rotation"]),
                "description": _text(row["description"]),
                "event_description": _text(row["event_description"]),
                "instagram_id": _text(row["instagram_id"]),
                "entrance_fee": row["entrance_fee"],
                "has_reusable_container": _bool(
                    row["has_reusable_"],
                    formula_booths[key]["has_reusable_"],
                    f"{key} has_reusable_container",
                ),
                "operations": booth_operations,
                "menus": booth_menus,
            }
            booth.update(BOOTH_OVERRIDES.get((name, zone), {}))
            booths.append(booth)
        if operations or menus:
            unknown = sorted({*operations, *menus})
            raise CommandError(f"부스 시트에 없는 임시키가 있습니다: {unknown}")

        output = Path(options["output"])
        output.parent.mkdir(parents=True, exist_ok=True)
        payload = {"source": xlsx_path.name, "booths": booths}
        content = json.dumps(payload, ensure_ascii=False, indent=2) + "\n"
        output.write_text(content, encoding="utf-8", newline="\n")

        self.stdout.write(
            self.style.SUCCESS(
                f"{output} 저장 — 부스 {len(booths)} · "
                f"운영일정 {sum(len(b['operations']) for b in booths)} · "
                f"메뉴 {sum(len(b['menus']) for b in booths)} (가격 미수령 {skipped_menus}개 제외)"
            )
        )
        if skipped_booths:
            self.stdout.write(
                self.style.WARNING(
                    f"좌표가 없어 제외한 장소 {len(skipped_booths)}곳: {', '.join(skipped_booths)}"
                )
            )
